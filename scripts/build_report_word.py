#!/usr/bin/env python3

from copy import deepcopy
import json
import os
import re
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
from zipfile import ZipFile

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'tools/python-docx/usr/lib/python3/dist-packages'))
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK,WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree

NAVY='153047'; TEAL='087E8B'


def shade(cell,color):
    el=OxmlElement('w:shd'); el.set(qn('w:fill'),color); cell._tc.get_or_add_tcPr().append(el)


def hyperlink(p,text,url):
    relation=p.part.relate_to(url,RT.HYPERLINK,is_external=True)
    el=OxmlElement('w:hyperlink'); el.set(qn('r:id'),relation)
    r=OxmlElement('w:r'); pr=OxmlElement('w:rPr')
    c=OxmlElement('w:color'); c.set(qn('w:val'),TEAL); pr.append(c)
    size=OxmlElement('w:sz'); size.set(qn('w:val'),'19'); pr.append(size); r.append(pr)
    t=OxmlElement('w:t'); t.text=text; r.append(t); el.append(r); p._p.append(el)


def set_font(run,size=10,color=NAVY,bold=False):
    run.font.name='DejaVu Sans'; run.font.size=Pt(size)
    run.font.color.rgb=RGBColor.from_string(color); run.bold=bold


def build():
    blocks=json.loads((ROOT/'docs/report_content.json').read_text())
    pandoc=shutil.which('pandoc') or str(ROOT/'tools/pandoc/usr/bin/pandoc')
    equations=[b for b in blocks if b['kind']=='equation']
    with tempfile.TemporaryDirectory(prefix='rov-equations-') as folder:
        folder=Path(folder)
        def tex(value):
            return re.sub(r'\{\\rm\s+([^}]+)\}', r'{\\mathrm{\1}}', value)
        (folder/'equations.md').write_text('\n\n'.join('$$\n'+tex(b['math'])+'\n$$' for b in equations))
        command=[pandoc,str(folder/'equations.md'),'-o',str(folder/'equations.docx')]
        data=ROOT/'tools/pandoc/usr/share/pandoc/data'
        env=os.environ.copy()
        if data.exists():
            command.extend(['--data-dir',str(data)])
            env['pandoc_datadir']=str(data.parent)
        subprocess.run(command,check=True,env=env)
        with ZipFile(folder/'equations.docx') as archive:
            xml=etree.fromstring(archive.read('word/document.xml'))
        maths=xml.findall('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath')
        assert len(maths)==len(equations),(len(maths),len(equations))
        maths=iter(maths)
        doc=Document(); sec=doc.sections[0]
        sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
        sec.top_margin=Inches(.94); sec.bottom_margin=Inches(.72)
        sec.left_margin=sec.right_margin=Inches(.66)
        sec.header_distance=sec.footer_distance=Inches(.38)
        normal=doc.styles['Normal']; normal.font.name='DejaVu Sans'; normal.font.size=Pt(10)
        normal.font.color.rgb=RGBColor.from_string(NAVY)
        normal.paragraph_format.line_spacing=Pt(13.4); normal.paragraph_format.space_after=Pt(7)
        normal.paragraph_format.widow_control=True
        for name,size in [('Heading 1',17),('Heading 2',11)]:
            st=doc.styles[name]; st.font.name='DejaVu Sans'; st.font.size=Pt(size)
            st.font.color.rgb=RGBColor.from_string(NAVY if name=='Heading 1' else TEAL)
            st.paragraph_format.space_before=Pt(0); st.paragraph_format.space_after=Pt(14 if name=='Heading 1' else 8)
        head=sec.header.paragraphs[0]; set_font(head.add_run('LTU–RAI  /  UNDERWATER NAVIGATION'),8.5,TEAL,True)
        foot=sec.footer.paragraphs[0]; foot.style=normal
        foot.paragraph_format.tab_stops.add_tab_stop(Inches(6.90),WD_TAB_ALIGNMENT.RIGHT)
        set_font(foot.add_run('Yecheng Wang  |  BlueROV2 survey  |  21 October 2026\t'),7.5,'647582')
        field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); foot._p.append(field)
        doc.core_properties.title='Sonar-guided underwater survey with BlueROV2'
        doc.core_properties.author='Yecheng Wang'
        doc.core_properties.subject='LTU-RAI PhD Recruitment Assignment'
        first=True
        for block in blocks:
            kind=block['kind']
            if kind=='page':
                p=doc.add_paragraph(block['title'],'Heading 1')
                p.paragraph_format.page_break_before=not first; first=False
            elif kind=='paragraph':
                p=doc.add_paragraph(); size=block['fontsize']
                set_font(p.add_run(block['text']),size)
                p.paragraph_format.line_spacing=Pt(13.4 if size>=10 else 12)
                if block['text'].startswith('Figure '):
                    p.paragraph_format.space_after=Pt(8)
            elif kind=='subheading':
                doc.add_paragraph(block['text'],'Heading 2')
            elif kind=='equation':
                p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(10)
                p.paragraph_format.line_spacing=1.15
                p.paragraph_format.tab_stops.add_tab_stop(Inches(3.30),WD_TAB_ALIGNMENT.CENTER)
                p.paragraph_format.tab_stops.add_tab_stop(Inches(6.92),WD_TAB_ALIGNMENT.RIGHT)
                p.add_run('\t'); p._p.append(deepcopy(next(maths))); set_font(p.add_run('\t('+str(block['number'])+')'),10)
            elif kind=='table':
                p=doc.add_paragraph(); set_font(p.add_run(block['caption']),9,TEAL,True)
                p.paragraph_format.keep_with_next=True; p.paragraph_format.space_after=Pt(8)
                table=doc.add_table(rows=1,cols=len(block['columns']))
                table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
                widths=[6.95*w for w in block['widths']]
                for col,width in zip(table.columns,widths): col.width=Inches(width)
                for i,values in enumerate([block['columns']]+block['rows']):
                    row=table.rows[0] if i==0 else table.add_row()
                    trPr=row._tr.get_or_add_trPr(); trPr.append(OxmlElement('w:cantSplit'))
                    if i==0: trPr.append(OxmlElement('w:tblHeader'))
                    for cell,width,value in zip(row.cells,widths,values):
                        cell.width=Inches(width); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        shade(cell,TEAL if i==0 else ('F0F5F7' if i%2 else 'FFFFFF'))
                        p=cell.paragraphs[0]; p.paragraph_format.space_before=Pt(4)
                        p.paragraph_format.space_after=Pt(4); p.paragraph_format.line_spacing=Pt(10.5)
                        set_font(p.add_run(str(value)),block['fontsize'],'FFFFFF' if i==0 else NAVY,i==0)
                gap=doc.add_paragraph(); gap.paragraph_format.line_spacing=Pt(1); gap.paragraph_format.space_after=Pt(6)
                set_font(gap.add_run(''),1)
            elif kind=='code':
                p=doc.add_paragraph(); p.paragraph_format.keep_together=True
                p.paragraph_format.line_spacing=Pt(9.5); p.paragraph_format.space_after=Pt(10)
                shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'F0F4F6'); p._p.get_or_add_pPr().append(shd)
                run=p.add_run(block['text']); run.font.name='DejaVu Sans Mono'; run.font.size=Pt(block['fontsize'])
            elif kind=='figure':
                p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.line_spacing=1.0
                p.paragraph_format.keep_with_next=True; p.paragraph_format.space_after=Pt(9)
                p.add_run().add_picture(block['path'],width=Inches(min(6.95,block['width_inches'])))
            elif kind=='link':
                p=doc.add_paragraph(); p.paragraph_format.keep_with_next=True; p.paragraph_format.space_after=Pt(4)
                hyperlink(p,block['text'],block['url'])
        doc.save(ROOT/'docs/report.docx')
    print(ROOT/'docs/report.docx')


if __name__=='__main__': build()
